"""content/slides.yaml + assets 그림 → output/*.pptx 조립 (python-pptx)."""
import sys
from pathlib import Path

import yaml
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent / "common"))
import style as S  # noqa: E402
import anim  # noqa: E402

FLOW_DIR = S.ROOT / "assets" / "flow"      # Google Flow 로 만든 분위기 이미지 (없으면 기본 디자인)
VIDEO_DIR = S.ROOT / "assets" / "videos"   # manim 영상 (mp4 + 포스터 png)

SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.67
CONTENT_W = SLIDE_W - 2 * MARGIN

COLOR_KEYS = {"volt": S.VOLT, "current": S.CURRENT, "resist": S.RESIST,
              "safe": S.SAFE, "spark": S.SPARK, "ink": S.INK}
# SPARK(노랑)는 흰 배경 위 글자색으로 쓰기엔 대비가 낮아서 글자에는 진한 색을 대신 쓴다
TEXT_ON_WHITE = {**COLOR_KEYS, "spark": "B07800"}


def hexrgb(h):
    return RGBColor.from_string(h.lstrip("#"))


# ── 기본 도형 ──────────────────────────────────────
def set_font(run, size, color=S.INK, bold=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = hexrgb(color)
    f.name = S.FONT_PPT
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):  # 한글(동아시아) 글꼴도 지정해야 맑은 고딕이 적용된다
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", S.FONT_PPT)


def add_text(slide, x, y, w, h, text, size=16, color=S.INK, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, name=None, line_spacing=1.1):
    """text 가 list 이면 문단 여러 개. 문단 항목이 (문자열, dict) 이면 그 문단만 스타일 변경."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = text if isinstance(text, list) else text.split("\n")  # "\n" = 줄바꿈 위치 지정
    for i, item in enumerate(paras):
        s, opt = (item if isinstance(item, tuple) else (item, {}))
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opt.get("align", align)
        p.line_spacing = line_spacing
        if i and opt.get("space_before"):
            p.space_before = Pt(opt["space_before"])
        r = p.add_run()
        r.text = s
        set_font(r, opt.get("size", size), opt.get("color", color), opt.get("bold", bold))
    return tb


def add_shape(slide, kind, x, y, w, h, fill, radius=None, name=None):
    sp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    style = sp._element.find(qn("p:style"))  # 테마 기본 그림자·테두리 제거
    if style is not None:
        sp._element.remove(style)
    sp.fill.solid()
    sp.fill.fore_color.rgb = hexrgb(fill)
    sp.line.fill.background()
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = radius / min(w, h)
    if name:
        sp.name = name
    return sp


def card(slide, x, y, w, h, fill=S.PANEL, name=None):
    return add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, radius=0.18, name=name)


def circle_label(slide, cx, cy, d, fill, text, size, color="FFFFFF", name=None):
    c = add_shape(slide, MSO_SHAPE.OVAL, cx - d / 2, cy - d / 2, d, d, fill, name=name)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    set_font(r, size, color, bold=True)
    return c


def add_image(slide, path, x, y, w=None, h=None, name=None):
    if not path.exists():
        print(f"  ! 그림 없음: {path.name} (자리표시 상자로 대체)")
        box = card(slide, x, y, w or 3, h or 3, fill=S.LINE)
        add_text(slide, x, y, w or 3, h or 3, f"[{path.stem}]", 14, S.INK_SOFT,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        return box
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y),
                                   Inches(w) if w else None, Inches(h) if h else None)
    if name:
        pic.name = name
    return pic


def set_alpha(shape, opacity):
    """도형 채우기 불투명도 (0~1)."""
    clr = shape.fill._xPr.find(qn("a:solidFill"))[0]
    etree.SubElement(clr, qn("a:alpha")).set("val", str(int(opacity * 100000)))


def flow_image(name):
    for ext in (".png", ".jpg", ".jpeg", ".webp"):
        p = FLOW_DIR / f"{name}{ext}"
        if p.exists():
            return p
    return None


def full_bleed(slide, path, overlay=(0.88, 0.15)):
    """슬라이드 전체를 채우는 배경 사진 (16:9 로 가운데 잘라냄) + 어두운 덮개."""
    pic = slide.shapes.add_picture(str(path), 0, 0, Inches(SLIDE_W), Inches(SLIDE_H))
    pic.name = "flow_bg"
    iw, ih = Image.open(path).size
    target = SLIDE_W / SLIDE_H
    if iw / ih > target:  # 옆으로 넓다 → 좌우를 자름
        cut = (1 - target * ih / iw) / 2
        pic.crop_left = pic.crop_right = cut
    else:
        cut = (1 - iw / ih / target) / 2
        pic.crop_top = pic.crop_bottom = cut
    # 왼쪽(글자 자리)은 진하게, 오른쪽(사진 주인공)은 옅게 덮는 가로 그라데이션
    shade = add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H, S.NIGHT, name="shade")
    shade.fill.gradient()
    shade.fill.gradient_angle = 0
    for stop, pos, opacity in zip(shade.fill.gradient_stops, (0, 100000), (overlay[0], overlay[1])):
        stop.color.rgb = hexrgb(S.NIGHT)
        clr = stop._gs[0]
        etree.SubElement(clr, qn("a:alpha")).set("val", str(int(opacity * 100000)))


def video_duration_ms(path):
    import av  # manim 설치 시 함께 설치됨
    with av.open(str(path)) as c:
        return int(c.duration / 1000) if c.duration else 10000


def add_video(slide, name, x, y, w):
    """manim 영상. 없으면 None."""
    mp4, poster = VIDEO_DIR / f"{name}.mp4", VIDEO_DIR / f"{name}.png"
    if not mp4.exists():
        print(f"  ! 영상 없음: {mp4.name} (그림으로 대체)")
        return None
    h = w * S.FIG_FULL[1] / S.FIG_FULL[0]
    mv = slide.shapes.add_movie(str(mp4), Inches(x), Inches(y), Inches(w), Inches(h),
                                poster_frame_image=str(poster) if poster.exists() else None,
                                mime_type="video/mp4")
    mv.name = name
    return mv.shape_id, video_duration_ms(mp4)


def fig(name):
    return S.FIG_DIR / f"{name}.png"


def icon(name):
    return S.ICON_DIR / f"{name}.png"


def background(slide, color):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = hexrgb(color)


def page_number(slide, n):
    add_text(slide, SLIDE_W - MARGIN - 1.0, SLIDE_H - 0.45, 1.0, 0.3, str(n), 11, S.INK_SOFT,
             align=PP_ALIGN.RIGHT, name="page")


def title_block(slide, d):
    add_text(slide, MARGIN, 0.42, CONTENT_W, 0.75, d["title"], 32, S.INK, bold=True,
             anchor=MSO_ANCHOR.MIDDLE, name="title")
    if d.get("lead"):
        add_text(slide, MARGIN, 1.2, CONTENT_W, 0.5, d["lead"], 18, S.INK_SOFT, name="lead")


def points_column(slide, points, x, y_top, w, y_bottom):
    """오른쪽 포인트 카드 (색 점 + 굵은 제목 + 설명)."""
    n = len(points)
    gap = 0.25
    h = min(1.45, (y_bottom - y_top - gap * (n - 1)) / n)
    y = y_top + (y_bottom - y_top - (h * n + gap * (n - 1))) / 2
    steps = []
    for i, pt in enumerate(points):
        col = COLOR_KEYS.get(pt.get("color", "ink"), S.INK)
        g = slide.shapes.add_group_shape()  # 카드 하나 = 그룹 하나 = 클릭 한 번
        g.name = f"point{i + 1}"
        card(g, x, y, w, h)
        add_shape(g, MSO_SHAPE.OVAL, x + 0.28, y + 0.3, 0.22, 0.22, col)
        add_text(g, x + 0.65, y + 0.2, w - 0.85, 0.42, pt["head"], 19,
                 TEXT_ON_WHITE.get(pt.get("color", "ink"), S.INK), bold=True, anchor=MSO_ANCHOR.MIDDLE)
        add_text(g, x + 0.65, y + 0.68, w - 0.85, h - 0.8, pt["body"], 15, S.INK)
        steps.append([g.shape_id])
        y += h + gap
    return steps


# ── 레이아웃별 ─────────────────────────────────────
def build_cover(slide, d, end=False):
    background(slide, S.NIGHT)
    bg = flow_image("flow_end" if end else "flow_cover")
    if bg:
        full_bleed(slide, bg)
    else:
        add_image(slide, fig(d["figure"]), 6.7, 0.65, w=6.2, name="hero")
    add_text(slide, 0.9, 2.2, 6.2, 2.2, d["title"], 50 if not end else 56, "FFFFFF", bold=True,
             anchor=MSO_ANCHOR.BOTTOM, name="title", line_spacing=1.0)
    add_text(slide, 0.9, 4.65, 6.0, 0.6, d["subtitle"], 22, S.SPARK, name="subtitle")
    if not end:
        add_text(slide, 0.9, 6.55, 6.0, 0.4, "그림과 비유로 배우는 전기의 기본", 14, "AEB6CC")


def build_section(slide, d):
    background(slide, S.NIGHT)
    bg = flow_image(f"flow_sec{d['num']}")
    if bg:
        full_bleed(slide, bg)
    add_text(slide, 0.9, 1.55, 4.0, 1.5, d["num"], 96, S.SPARK, bold=True,
             anchor=MSO_ANCHOR.BOTTOM, name="num", line_spacing=0.9)
    add_text(slide, 0.9, 3.2, 7.6, 1.0, d["title"], 44, "FFFFFF", bold=True, name="title")
    add_text(slide, 0.9, 4.35, 7.6, 0.6, d["subtitle"], 20, "C9D0E2", name="subtitle")
    if not bg:
        add_image(slide, icon(d["icon"]), 9.1, 2.25, w=3.0, name="icon")


def build_full(slide, d):
    title_block(slide, d)
    video = add_video(slide, d["video"], MARGIN, 2.05, S.FIG_FULL[0]) if d.get("video") else None
    if not video:
        add_image(slide, fig(d.get("figure") or d["video"]), MARGIN, 2.05, w=S.FIG_FULL[0], name="figure")
    return {"video": video}


def build_split(slide, d):
    title_block(slide, d)
    fh = 5.1
    fw = fh * S.FIG_SPLIT[0] / S.FIG_SPLIT[1]
    add_image(slide, fig(d["figure"]), 0.55, 1.95, w=fw, h=fh, name="figure")
    px = 0.55 + fw + 0.4
    return {"steps": points_column(slide, d["points"], px, 2.05, SLIDE_W - MARGIN - px, 6.95)}


def build_cards(slide, d):
    title_block(slide, d)
    cards = d["cards"]
    n = len(cards)
    gap = 0.3
    steps = []
    if n <= 5:  # 한 줄, 세로형 카드
        w = (CONTENT_W - gap * (n - 1)) / n
        y, h = 2.15, 4.55
        isz = 1.6 if n <= 3 else 1.35
        for i, c in enumerate(cards):
            x = MARGIN + i * (w + gap)
            col = TEXT_ON_WHITE.get(c.get("color", "ink"), S.INK)
            g = slide.shapes.add_group_shape()
            g.name = f"card{i + 1}"
            card(g, x, y, w, h)
            add_image(g, icon(c["icon"]), x + (w - isz) / 2, y + 0.45, w=isz, h=isz)
            ty = y + 0.45 + isz + 0.35
            hs = 24 if n <= 3 else 20
            add_text(g, x + 0.2, ty, w - 0.4, 0.55, c["head"], hs, col, bold=True, align=PP_ALIGN.CENTER)
            add_text(g, x + 0.2, ty + 0.65, w - 0.4, h - (ty - y) - 0.8, c["body"],
                     17 if n <= 3 else 15, S.INK, align=PP_ALIGN.CENTER)
            steps.append([g.shape_id])
    else:  # 2줄 x 3칸, 가로형 카드
        cols = 3
        w = (CONTENT_W - gap * (cols - 1)) / cols
        h = 2.15
        for i, c in enumerate(cards):
            x = MARGIN + (i % cols) * (w + gap)
            y = 2.1 + (i // cols) * (h + gap)
            col = TEXT_ON_WHITE.get(c.get("color", "ink"), S.INK)
            g = slide.shapes.add_group_shape()
            g.name = f"card{i + 1}"
            card(g, x, y, w, h)
            add_image(g, icon(c["icon"]), x + 0.25, y + (h - 1.1) / 2, w=1.1, h=1.1)
            # 제목·설명을 한 상자에 넣어 제목이 두 줄이 돼도 겹치지 않게 한다
            body = c["body"] if isinstance(c["body"], list) else [c["body"]]
            add_text(g, x + 1.55, y + 0.2, w - 1.75, h - 0.4,
                     [(c["head"], {"size": 20, "color": col, "bold": True})]
                     + [(b, {"size": 15, "space_before": 6 if k == 0 else 0}) for k, b in enumerate(body)],
                     anchor=MSO_ANCHOR.MIDDLE)
            steps.append([g.shape_id])
    return {"steps": steps}


def build_chart(slide, d):
    title_block(slide, d)
    rows = sorted(d["chart"]["data"], key=lambda r: r[1])  # 큰 값이 위로 오도록
    cd = CategoryChartData()
    cd.categories = [r[0] for r in rows]
    cd.add_series("소비전력", [r[1] for r in rows])
    x, y, w, h = MARGIN, 2.0, 7.4, 5.0
    gframe = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    gframe.name = "chart"
    ch = gframe.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.size = Pt(14)
    ch.font.name = S.FONT_PPT
    ch.font.color.rgb = hexrgb(S.INK)
    plot = ch.plots[0]
    plot.gap_width = 45
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = '#,##0"W"'
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(14)
    dl.font.bold = True
    dl.font.color.rgb = hexrgb(S.INK)
    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = False
    va.maximum_scale = max(r[1] for r in rows) * 1.22
    ca = ch.category_axis
    ca.format.line.fill.background()
    ca.tick_labels.font.size = Pt(15)
    ser = plot.series[0]
    top = {r[0] for r in sorted(rows, key=lambda r: -r[1])[:3]}
    for i, r in enumerate(rows):
        f = ser.points[i].format.fill
        f.solid()
        f.fore_color.rgb = hexrgb(S.RESIST if r[0] in top else (S.SAFE if r[1] <= 20 else "B9C1D1"))
    steps = points_column(slide, d["points"], x + w + 0.4, 2.2, SLIDE_W - MARGIN - (x + w + 0.4), 6.6)
    add_text(slide, x, 7.0, w, 0.3, "※ 제품마다 다른 대략적인 값", 11, S.INK_SOFT)
    return {"steps": steps}


def build_quiz(slide, d):
    add_text(slide, MARGIN, 0.42, CONTENT_W, 0.75, d["title"], 32, S.INK, bold=True,
             anchor=MSO_ANCHOR.MIDDLE, name="title")
    reveal = d.get("reveal")
    items = d["items"]
    gap, h = 0.22, 1.12
    y = 1.65
    steps = []
    for i, it in enumerate(items):
        card(slide, MARGIN, y, CONTENT_W, h, name=f"q{i + 1}")
        circle_label(slide, MARGIN + 0.6, y + h / 2, 0.62, S.INK, str(i + 1), 20)
        if reveal:  # 문제는 보이고, 클릭할 때마다 정답(O/X + 이유)이 하나씩 나타남
            add_text(slide, MARGIN + 1.2, y + 0.16, CONTENT_W - 3.0, 0.45, it["q"], 21, S.INK, bold=True)
            ok = it["a"] == "O"
            g = slide.shapes.add_group_shape()
            g.name = f"a{i + 1}"
            add_text(g, MARGIN + 1.2, y + 0.64, CONTENT_W - 3.0, 0.35, it["why"], 15, S.INK_SOFT)
            circle_label(g, SLIDE_W - MARGIN - 0.85, y + h / 2, 0.86, S.SAFE if ok else S.RESIST,
                         "O" if ok else "X", 30)
            steps.append([g.shape_id])
        else:
            add_text(slide, MARGIN + 1.2, y, CONTENT_W - 3.6, h, it["q"], 22, S.INK, bold=True,
                     anchor=MSO_ANCHOR.MIDDLE)
            for k, (lab, col) in enumerate((("O", S.SAFE), ("X", S.RESIST))):
                cx = SLIDE_W - MARGIN - 1.55 + k * 1.0
                c = circle_label(slide, cx, y + h / 2, 0.78, "FFFFFF", lab, 26, col)
                c.line.color.rgb = hexrgb(col)
                c.line.width = Pt(2.5)
        y += h + gap
    return {"steps": steps}


BUILDERS = {
    "cover": build_cover,
    "end": lambda s, d: build_cover(s, d, end=True),
    "section": build_section,
    "full": build_full,
    "split": build_split,
    "cards": build_cards,
    "chart": build_chart,
    "quiz": build_quiz,
}


def main():
    data = yaml.safe_load((S.CONTENT_DIR / "slides.yaml").read_text(encoding="utf-8"))
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    prs.core_properties.title = data["meta"]["title"]
    blank = prs.slide_layouts[6]
    for n, d in enumerate(data["slides"], start=1):
        slide = prs.slides.add_slide(blank)
        fx = BUILDERS[d["layout"]](slide, d) or {}
        if not d.get("static"):  # static: true 이면 애니메이션 없이
            anim.set_timing(slide, fx.get("steps", []), fx.get("video"))
        if d["layout"] not in ("cover", "end", "section"):
            page_number(slide, n)
        if d.get("notes"):
            slide.notes_slide.notes_text_frame.text = d["notes"].strip()
    S.OUTPUT_DIR.mkdir(exist_ok=True)
    out = S.OUTPUT_DIR / data["meta"]["file"]
    prs.save(out)
    print(f"슬라이드 {len(data['slides'])}장 → output/{out.name}")


if __name__ == "__main__":
    main()
